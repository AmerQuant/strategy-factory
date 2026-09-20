"""Moneta MT5 ECN broker file -> normalized table -> cost profiles and symbol map (T06b).

Two steps (D-317, D-340):

1. :func:`import_spec` reads the broker xlsx (openpyxl, dev-only dependency, imported lazily)
   into the committed, normalized ``configs/costs/moneta/moneta_spec.csv``. A sidecar
   ``moneta_spec.csv.meta.json`` records the source path, its SHA-256 (checked against the
   raw-store manifest), the file date and the import time. Re-importing the same file gives
   a byte-identical CSV. The runtime never reads the xlsx.
2. :func:`build` maps broker symbols to research symbols (D-323, D-324, D-325) and writes the
   generated ``moneta_profiles.yaml``, ``symbol_map.csv``, ``assignments.yaml`` and the
   mapping review CSV for the user. A broker symbol whose ticker matches a research symbol but
   whose **name belongs to another company is never ticker-mapped** (D-341, e.g. broker ``ESL``
   = Estee Lauder vs research ``ESL``): only a manual override in ``symbol_overrides.csv`` can
   map it. Names come from ``mapping.yaml: names_file``; after the Alpaca asset list (ETFs
   included) is downloaded with ``scripts/download_alpaca_assets.ps1`` (the user runs it,
   D-031), point that file at it and rebuild for a reduced review list.

Units in the normalized table: ``point_size`` = 10^-digits; ``point_value`` = money per point
per lot in ``quote_ccy``; ``contract_size`` = instrument units per lot; volumes in lots;
``spread_points`` as in the file and ``spread_price`` = points x point size (full spread);
commission and swap amounts exactly as in the file (their meaning is applied in ``build``).
"""

from __future__ import annotations

import csv
import datetime as dt
import difflib
import io
import json
import math
import re
import statistics
from collections import Counter
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import asdict, dataclass, fields, replace
from pathlib import Path
from typing import Any, Literal

import yaml
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from strategy_factory.core.errors import ConfigError, DataError
from strategy_factory.costs.profile import (
    ASSIGNMENTS_FILE,
    MONETA_DIR,
    MONETA_PROFILES_FILE,
    CostProfile,
    Slippage,
    Weekday,
)

MONETA_CONFIG = "moneta.yaml"
MAPPING_CONFIG = "mapping.yaml"
SPEC_CSV = "moneta_spec.csv"
SPEC_META = "moneta_spec.csv.meta.json"
SYMBOL_MAP_CSV = "symbol_map.csv"
DUKASCOPY_MAP_CSV = "dukascopy_map.csv"
OVERRIDES_CSV = "symbol_overrides.csv"
UNMAPPABLE = "UNMAPPABLE"
SHEETS = ("Forex&Metals", "Commodities", "Indices", "Share_CFDs", "Crypto", "Bond_CFDs")
WEEKDAYS: dict[str, Weekday] = {
    "monday": "MON", "tuesday": "TUE", "wednesday": "WED", "thursday": "THU",
    "friday": "FRI", "saturday": "SAT", "sunday": "SUN",
}  # fmt: skip
RowStatus = Literal["ok", "incomplete", "shifted", "unparsed"]


class _Frozen(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")


# -- configuration ---------------------------------------------------------------------------
class SourceCfg(_Frozen):
    spec_file: str
    file_date: str


class CommissionCfg(_Frozen):
    per_lot_sides: int = Field(gt=0)
    per_trade_sides: int = Field(gt=0)


class AssumedSwap(_Frozen):
    long: float
    short: float
    triple_weekday: Weekday


class SwapCfg(_Frozen):
    percent_day_count: int = Field(gt=0)
    rollover_time_local: str
    rollover_tz: str
    rollover_weekdays: tuple[Weekday, ...]
    etf_assumed: AssumedSwap


class SpreadCfg(_Frozen):
    data_classes: tuple[str, ...]


class MonetaConfig(_Frozen):
    source: SourceCfg
    commission: CommissionCfg
    swap: SwapCfg
    spread: SpreadCfg
    pip_per_point: float = Field(gt=0)
    pip_classes: tuple[str, ...]  # asset classes whose slippage is given in pips
    us_share_region: str
    etf_region: str
    proxy_profile: str
    proxy_asset_class: str  # research asset class of the D-324 proxy (its slippage key)
    slippage: dict[str, Slippage]


class MappingConfig(_Frozen):
    names_file: str
    broker_regions: tuple[str, ...]
    min_name_score: float = Field(ge=0, le=1)
    candidate_min_score: float = Field(ge=0, le=1)
    prices_dir: str | None = None  # raw daily files for the review price check (D-341)


def _yaml(path: Path) -> dict[str, Any]:
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except (OSError, yaml.YAMLError) as exc:
        raise ConfigError(f"cannot read config: {exc}", config_path=path) from exc
    if not isinstance(data, dict):
        raise ConfigError("config root must be a mapping", config_path=path)
    return data


def load_moneta_config(moneta_dir: Path) -> MonetaConfig:
    path = moneta_dir / MONETA_CONFIG
    try:
        return MonetaConfig.model_validate(_yaml(path))
    except ValidationError as exc:
        raise ConfigError(f"invalid Moneta config: {exc}", config_path=path) from exc


def load_mapping_config(moneta_dir: Path) -> MappingConfig:
    path = moneta_dir / MAPPING_CONFIG
    try:
        return MappingConfig.model_validate(_yaml(path))
    except ValidationError as exc:
        raise ConfigError(f"invalid mapping config: {exc}", config_path=path) from exc


# -- the normalized table ----------------------------------------------------------------------
@dataclass(frozen=True)
class SpecRow:
    sheet: str
    row: int  # spreadsheet row number (1 = header)
    region: str
    broker_symbol: str
    description: str
    digits: int | None
    point_size: float | None
    contract_size: float | None
    contract_unit: str
    quote_ccy: str
    point_value: float | None
    min_volume_lots: float | None
    volume_step_lots: float | None
    spread_points: float | None
    spread_price: float | None
    commission_model: str  # none | per_lot | per_trade | percent_per_lot
    commission_amount: float | None
    commission_ccy: str
    swap_model: str  # none | points | currency | percent
    swap_long: float | None
    swap_short: float | None
    triple_weekday: str
    quote_sample: float | None
    trading_time_server: str
    row_status: RowStatus
    status_reason: str


SPEC_COLUMNS = tuple(f.name for f in fields(SpecRow))
_INT_COLS = {"row", "digits"}
_FLOAT_COLS = {
    "point_size", "contract_size", "point_value", "min_volume_lots", "volume_step_lots",
    "spread_points", "spread_price", "commission_amount", "swap_long", "swap_short",
    "quote_sample",
}  # fmt: skip


def _fmt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return repr(v)
    return str(v)


def write_spec_csv(rows: Sequence[SpecRow], path: Path) -> None:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(SPEC_COLUMNS)
    for r in rows:
        w.writerow([_fmt(getattr(r, c)) for c in SPEC_COLUMNS])
    path.write_bytes(buf.getvalue().encode("utf-8"))


def read_spec_csv(path: Path) -> list[SpecRow]:
    if not path.is_file():
        raise ConfigError(
            "normalized broker table not found (sfac costs moneta import)", config_path=path
        )
    out: list[SpecRow] = []
    with path.open(encoding="utf-8", newline="") as fh:
        for d in csv.DictReader(fh):
            vals: dict[str, Any] = {}
            for c in SPEC_COLUMNS:
                raw = d[c]
                if c in _INT_COLS:
                    vals[c] = int(raw) if raw else None
                elif c in _FLOAT_COLS:
                    vals[c] = float(raw) if raw else None
                else:
                    vals[c] = raw
            out.append(SpecRow(**vals))
    return out


# -- parsing one spreadsheet row ---------------------------------------------------------------
_MISSING = {"", "-"}


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() in _MISSING)


def _num(v: Any) -> float | None:
    if _blank(v):
        return None
    if isinstance(v, bool):
        raise ValueError(f"not a number: {v!r}")
    if isinstance(v, int | float):
        return float(v)
    return float(str(v).strip())


def _amount_unit(v: Any) -> tuple[float, str]:
    m = re.fullmatch(r"\s*([0-9.]+(?:[eE][-+]?[0-9]+)?)\s+(\S.*?)\s*", str(v))
    if not m:
        raise ValueError(f"expected '<number> <unit>', got {v!r}")
    return float(m.group(1)), m.group(2)


def _commission(v: Any) -> tuple[str, float | None, str]:
    if _blank(v):
        return "none", None, ""
    text = str(v).strip()
    patterns = (
        (r"([0-9.]+) ([A-Z]{3}) per lot", "per_lot"),
        (r"([0-9.]+) ([A-Z]{3}) per trade", "per_trade"),
    )
    for pat, model in patterns:
        m = re.fullmatch(pat, text)
        if m:
            return model, float(m.group(1)), m.group(2)
    m = re.fullmatch(r"([0-9.]+) percentage per lot", text)
    if m:
        return "percent_per_lot", float(m.group(1)), ""
    raise ValueError(f"unknown commission {text!r}")


def _swap_model(v: Any) -> str:
    if _blank(v):
        return "none"
    models = {"in points": "points", "in currency": "currency", "in percentage terms": "percent"}
    text = str(v).strip()
    if text not in models:
        raise ValueError(f"unknown swap type {text!r}")
    return models[text]


def _weekday(v: Any) -> str:
    if _blank(v):
        return ""
    key = str(v).strip().lower()
    if key not in WEEKDAYS:
        raise ValueError(f"not a weekday: {v!r}")
    return WEEKDAYS[key]


def _is_weekday(v: Any) -> bool:
    return isinstance(v, str) and v.strip().lower() in WEEKDAYS


def _text(v: Any) -> str:
    if v is None:
        return ""
    return " ".join(str(v).replace("\xa0", " ").split())


def _trading_time(v: Any) -> str:
    if v is None:
        return ""
    return "; ".join(line.strip() for line in str(v).splitlines() if line.strip())


def normalize_row(sheet: str, row_no: int, cells: Mapping[str, Any]) -> SpecRow:
    """One spreadsheet row (header -> value) as a :class:`SpecRow`; never raises."""
    symbol = _text(cells.get("Symbol"))
    base: dict[str, Any] = {
        "sheet": sheet,
        "row": row_no,
        "region": _text(cells.get("Region")),
        "broker_symbol": symbol,
        "description": _text(cells.get("Description")),
        "trading_time_server": _trading_time(cells.get("Trading time")),
    }
    empty: dict[str, Any] = {
        "digits": None, "point_size": None, "contract_size": None, "contract_unit": "",
        "quote_ccy": "", "point_value": None, "min_volume_lots": None, "volume_step_lots": None,
        "spread_points": None, "spread_price": None, "commission_model": "",
        "commission_amount": None, "commission_ccy": "", "swap_model": "", "swap_long": None,
        "swap_short": None, "triple_weekday": "", "quote_sample": None,
    }  # fmt: skip
    if _is_weekday(cells.get("Commission")) or isinstance(cells.get("3-day swap"), int | float):
        return SpecRow(
            **base, **empty, row_status="shifted", status_reason="cells shifted by a column"
        )
    try:
        digits_f = _num(cells.get("Digits"))
        if digits_f is None or digits_f != int(digits_f):
            raise ValueError(f"invalid digits {cells.get('Digits')!r}")
        digits = int(digits_f)
        point_size = float(f"1e-{digits}")
        point_value, quote_ccy = _amount_unit(cells.get("Point value"))
        contract_size, contract_unit = _amount_unit(cells.get("Contract Size"))
        spread_points = _num(cells.get("Spread(For reference only)"))
        c_model, c_amount, c_ccy = _commission(cells.get("Commission"))
        s_model = _swap_model(cells.get("Swap Type"))
        s_long, s_short = _num(cells.get("SWAP long")), _num(cells.get("SWAP short"))
        triple = _weekday(cells.get("3-day swap"))
        row = SpecRow(
            **base,
            digits=digits,
            point_size=point_size,
            contract_size=contract_size,
            contract_unit=contract_unit,
            quote_ccy=quote_ccy,
            point_value=point_value,
            min_volume_lots=_num(cells.get("Min volume per click")),
            volume_step_lots=_num(cells.get("Volume Step")),
            spread_points=spread_points,
            spread_price=None if spread_points is None else spread_points * point_size,
            commission_model=c_model,
            commission_amount=c_amount,
            commission_ccy=c_ccy or (quote_ccy if c_model == "percent_per_lot" else ""),
            swap_model=s_model,
            swap_long=s_long,
            swap_short=s_short,
            triple_weekday=triple,
            quote_sample=_num(cells.get("Quote sample")),
            row_status="ok",
            status_reason="",
        )
    except ValueError as exc:
        return SpecRow(**base, **empty, row_status="unparsed", status_reason=str(exc))
    missing = [
        name
        for name, val in (
            ("spread", row.spread_points),
            ("quote sample", row.quote_sample),
            ("volume step", row.volume_step_lots),
            ("min volume", row.min_volume_lots),
        )
        if val is None
    ]
    if row.swap_model != "none" and (
        row.swap_long is None or row.swap_short is None or not row.triple_weekday
    ):
        missing.append("swap values or triple day")
    if missing:
        return SpecRow(
            **{
                **asdict(row),
                "row_status": "incomplete",
                "status_reason": "missing " + ", ".join(missing),
            }
        )
    return row


# -- import ------------------------------------------------------------------------------------
def read_workbook(path: Path) -> list[tuple[str, int, dict[str, Any]]]:
    """(sheet, row number, header -> value) for every data row of the six sheets."""
    try:
        import openpyxl  # dev-only dependency (D-317)
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise ConfigError(
            "openpyxl is not installed: `sfac costs moneta import` needs the dev dependencies "
            "(uv sync)"
        ) from exc
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        missing = [s for s in SHEETS if s not in wb.sheetnames]
        if missing:
            raise DataError(f"broker file lacks sheets {missing}")
        out: list[tuple[str, int, dict[str, Any]]] = []
        for sheet in SHEETS:
            rows = wb[sheet].iter_rows(values_only=True)
            header = [_text(h) for h in next(rows)]
            for i, values in enumerate(rows, start=2):
                if all(_blank(v) for v in values):
                    continue
                out.append((sheet, i, dict(zip(header, values, strict=False))))
        return out
    finally:
        wb.close()


@dataclass(frozen=True)
class ImportReport:
    source: Path
    sha256: str
    counts: dict[str, int]  # sheet -> rows
    status_counts: dict[str, int]
    not_ok: list[SpecRow]


def _sha256(path: Path) -> str:
    from strategy_factory.data.hashing import file_sha256

    return file_sha256(path)


def import_spec(xlsx: Path, moneta_dir: Path, *, source_label: str, file_date: str) -> ImportReport:
    """Parse ``xlsx`` into ``moneta_dir/moneta_spec.csv`` + sidecar (see module docstring)."""
    sha = _sha256(xlsx)
    manifest = xlsx.with_name(xlsx.name + ".manifest.json")
    if not manifest.is_file():
        raise DataError(f"broker file has no raw-store manifest ({manifest.name}); D-340")
    expected = json.loads(manifest.read_text(encoding="utf-8")).get("sha256")
    if expected != sha:
        raise DataError(f"broker file SHA-256 {sha} does not match its manifest ({expected})")
    rows = [normalize_row(s, i, cells) for s, i, cells in read_workbook(xlsx)]
    moneta_dir.mkdir(parents=True, exist_ok=True)
    write_spec_csv(rows, moneta_dir / SPEC_CSV)
    meta = {
        "source": source_label,
        "sha256": sha,
        "size_bytes": xlsx.stat().st_size,
        "file_date": file_date,
        "manifest_checked": True,
        "imported_at": dt.datetime.now(dt.UTC).replace(microsecond=0).isoformat(),
        "rows": len(rows),
    }
    (moneta_dir / SPEC_META).write_bytes((json.dumps(meta, indent=2) + "\n").encode("utf-8"))
    return ImportReport(
        source=xlsx,
        sha256=sha,
        counts=dict(Counter(r.sheet for r in rows)),
        status_counts=dict(Counter(r.row_status for r in rows)),
        not_ok=[r for r in rows if r.row_status != "ok"],
    )


def spec_sha256(moneta_dir: Path) -> str:
    path = moneta_dir / SPEC_META
    if not path.is_file():
        raise ConfigError(
            "broker table sidecar not found (sfac costs moneta import)", config_path=path
        )
    sha = json.loads(path.read_text(encoding="utf-8")).get("sha256")
    if not isinstance(sha, str) or len(sha) != 64:
        raise ConfigError("sidecar has no valid sha256", config_path=path)
    return sha


# -- mapping -----------------------------------------------------------------------------------
_STOP = frozenset(
    [
        "inc",
        "incorporated",
        "corp",
        "corporation",
        "co",
        "company",
        "ltd",
        "limited",
        "plc",
        "the",
        "class",
        "common",
        "stock",
        "shares",
        "share",
        "holdings",
        "holding",
        "group",
        "sa",
        "nv",
        "ag",
        "adr",
        "ads",
        "ordinary",
        "com",
        "llc",
        "lp",
        "trust",
        "se",
        "a",
        "b",
        "c",
        "new",
    ]
)


def _tokens(name: str) -> list[str]:
    text = name.lower().replace("&", " and ")
    return [t for t in re.split(r"[^a-z0-9]+", text) if t and t not in _STOP]


def name_score(a: str, b: str) -> float:
    """Similarity of two company names in [0, 1] after dropping legal-form words.

    The maximum of the character ratio of the joined tokens and the token containment
    ``|A & B| / min(|A|, |B|)``.
    """
    ta, tb = _tokens(a), _tokens(b)
    if not ta or not tb:
        return 0.0
    ratio = difflib.SequenceMatcher(None, " ".join(ta), " ".join(tb)).ratio()
    contain = len(set(ta) & set(tb)) / min(len(set(ta)), len(set(tb)))
    return max(ratio, contain)


@dataclass(frozen=True)
class MapRow:
    research_symbol: str
    broker_symbol: str
    method: str  # manual | ticker_exact | override
    name_score: float | None
    note: str


@dataclass(frozen=True)
class ReviewRow:
    broker_symbol: str
    region: str
    broker_description: str
    status: str  # pending_review | unmappable
    candidate_research_symbol: str
    candidate_name: str
    name_score: float | None
    reason: str
    # price check (D-341): the broker's quote sample against our own last close at the file
    # date. A ratio far from 1, or a stale close date, means the ticker is another instrument.
    broker_quote_sample: float | None = None
    research_close: float | None = None
    research_close_date: str = ""
    price_ratio: float | None = None


@dataclass(frozen=True)
class SymbolMapping:
    mapped: list[MapRow]
    review: list[ReviewRow]

    def by_research(self) -> dict[str, str]:
        return {m.research_symbol: m.broker_symbol for m in self.mapped}


def _read_csv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise ConfigError("file not found", config_path=path)
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def read_names(path: Path) -> dict[str, str]:
    return {r["symbol"]: r["name"] for r in _read_csv(path) if r.get("symbol")}


def map_symbols(
    spec: Sequence[SpecRow],
    research: Mapping[str, str],
    names: Mapping[str, str],
    dukascopy_map: Sequence[Mapping[str, str]],
    overrides: Sequence[Mapping[str, str]],
    cfg: MappingConfig,
) -> SymbolMapping:
    """Broker <-> research mapping (D-323, D-325); see ``docs/tasks/T06b_moneta_costs.md``."""
    by_broker = {r.broker_symbol: r for r in spec}
    mapped: list[MapRow] = []
    review: list[ReviewRow] = []
    for d in dukascopy_map:
        sym, bsym = d["research_symbol"], d["broker_symbol"]
        row = by_broker.get(bsym)
        if row is None or row.row_status != "ok":
            raise ConfigError(f"{DUKASCOPY_MAP_CSV}: broker symbol {bsym!r} missing or not ok")
        if sym not in research:
            raise ConfigError(f"{DUKASCOPY_MAP_CSV}: {sym!r} is not a research symbol")
        if bsym in {m.broker_symbol for m in mapped}:
            raise ConfigError(f"{DUKASCOPY_MAP_CSV}: broker symbol {bsym!r} listed twice")
        mapped.append(MapRow(sym, bsym, "manual", None, d.get("note", "")))

    over = {o["broker_symbol"]: o for o in overrides}
    unknown = sorted(set(over) - set(by_broker))
    if unknown:
        raise ConfigError(f"{OVERRIDES_CSV}: unknown broker symbols {unknown}")
    token_index: dict[str, set[str]] = {}
    for sym, nm in names.items():
        if sym in research:
            for t in set(_tokens(nm)):
                token_index.setdefault(t, set()).add(sym)

    def best(desc: str) -> tuple[str, float]:
        cands: set[str] = set()
        for t in set(_tokens(desc)):
            cands |= token_index.get(t, set())
        scored = sorted(
            ((name_score(desc, names[c]), c) for c in cands), key=lambda x: (-x[0], x[1])
        )
        return (scored[0][1], scored[0][0]) if scored else ("", 0.0)

    for r in spec:
        if r.region not in cfg.broker_regions:
            continue

        def rev(status: str, cand: str, score: float | None, reason: str, r: SpecRow = r) -> None:
            review.append(
                ReviewRow(
                    r.broker_symbol,
                    r.region,
                    r.description,
                    status,
                    cand,
                    names.get(cand, ""),
                    score,
                    reason,
                )
            )

        if r.row_status != "ok":
            if r.broker_symbol in over and over[r.broker_symbol]["research_symbol"] != UNMAPPABLE:
                raise ConfigError(  # D-350 (1): never a silent placeholder
                    f"{OVERRIDES_CSV}: {r.broker_symbol!r} maps to a {r.row_status} broker row "
                    f"({r.status_reason}); an incomplete row cannot yield a cost profile"
                )
            rev("unmappable", "", None, f"broker row {r.row_status}: {r.status_reason}")
            continue
        if r.broker_symbol in over:
            target = over[r.broker_symbol]["research_symbol"]
            note = over[r.broker_symbol].get("note", "")
            if target == UNMAPPABLE:
                rev("unmappable", "", None, f"override: {note}")
            elif target not in research:
                rev(
                    "unmappable",
                    target,
                    None,
                    f"override target not in the research universe ({note})",
                )
            else:
                sc = name_score(r.description, names[target]) if target in names else None
                mapped.append(MapRow(target, r.broker_symbol, "override", sc, note))
            continue
        if r.broker_symbol in research:
            rname = names.get(r.broker_symbol)
            if rname is None:
                rev(
                    "pending_review",
                    r.broker_symbol,
                    None,
                    "ticker match, no research name available",
                )
                continue
            sc = name_score(r.description, rname)
            if sc >= cfg.min_name_score:
                mapped.append(MapRow(r.broker_symbol, r.broker_symbol, "ticker_exact", sc, ""))
            else:  # another company with the same ticker: manual override only (D-341)
                rev("pending_review", r.broker_symbol, sc, "ticker match, name differs")
            continue
        cand, sc = best(r.description)
        if cand and sc >= cfg.candidate_min_score:
            rev("pending_review", cand, sc, "name match only (never auto-mapped)")
        else:
            rev("pending_review", cand, sc if cand else None, "no ticker or name match")

    dup = [s for s, n in Counter(m.research_symbol for m in mapped).items() if n > 1]
    if dup:
        raise ConfigError(f"research symbols mapped to two broker symbols: {sorted(dup)}")
    mapped.sort(key=lambda m: m.research_symbol)
    review.sort(key=lambda v: (v.status, v.broker_symbol))
    return SymbolMapping(mapped, review)


# -- profiles ----------------------------------------------------------------------------------
def with_price_check(
    review: Sequence[ReviewRow], spec: Sequence[SpecRow], day: dt.date, prices_root: Path
) -> list[ReviewRow]:
    """Add the quote-sample / our-close ratio to every candidate row (D-341)."""
    from strategy_factory.data.raw_prices import last_close_on_or_before

    samples = {r.broker_symbol: r.quote_sample for r in spec}
    out: list[ReviewRow] = []
    for row in review:
        sample = samples.get(row.broker_symbol)
        found = (
            last_close_on_or_before(row.candidate_research_symbol, day, prices_root)
            if row.candidate_research_symbol
            else None
        )
        close, close_day = found if found is not None else (None, None)
        ratio = sample / close if (sample and close) else None
        out.append(
            replace(
                row,
                broker_quote_sample=sample,
                research_close=close,
                research_close_date=close_day.isoformat() if close_day else "",
                price_ratio=ratio,
            )
        )
    return out


def profile_name(broker_symbol: str) -> str:
    return f"moneta_{broker_symbol}"


def _source_note(r: SpecRow, sha: str, cfg: MonetaConfig) -> str:
    return (
        f"Moneta MT5 ECN spec {Path(cfg.source.spec_file).name} sha256={sha} "
        f"file date {cfg.source.file_date}; sheet {r.sheet} row {r.row} ({r.broker_symbol})"
    )


def _rollover(cfg: MonetaConfig, triple: str) -> dict[str, Any]:
    return {
        "rollover_time_local": cfg.swap.rollover_time_local,
        "rollover_tz": cfg.swap.rollover_tz,
        "rollover_weekdays": list(cfg.swap.rollover_weekdays),
        "triple_weekday": triple,
    }


def _bps(r: SpecRow) -> float:
    assert r.spread_price is not None and r.quote_sample is not None
    return r.spread_price / r.quote_sample * 1e4


def profile_dict(r: SpecRow, asset_class: str, sha: str, cfg: MonetaConfig) -> dict[str, Any]:
    """The generated profile of one mapped broker row (D-318 ... D-322, D-520 ... D-523)."""
    if r.row_status != "ok":
        raise ConfigError(f"{r.broker_symbol}: broker row is {r.row_status}")
    assert r.point_size is not None and r.spread_price is not None
    note = _source_note(r, sha, cfg)
    to_verify: list[str] = []
    d: dict[str, Any] = {"name": profile_name(r.broker_symbol), "status": "verified"}
    if asset_class in cfg.spread.data_classes:
        spread: dict[str, Any] = {"mode": "broker_scaled", "broker_spread": r.spread_price}
    else:
        bps = _bps(r)
        spread = {"mode": "fixed", "fixed": {"value": bps, "unit": "bps"}}
        note += (
            f"; spread {r.spread_points:g} points x {r.point_size:g} / quote sample "
            f"{r.quote_sample:g} = {bps:.6g} bps"
        )
    if r.commission_model == "none":
        commission: dict[str, Any] = {"model": "none"}
    elif r.commission_model == "per_lot":
        assert r.commission_amount is not None and r.contract_size is not None
        commission = {
            "model": "per_lot",
            "lot_size": r.contract_size,
            "per_lot_per_side": r.commission_amount / cfg.commission.per_lot_sides,
            "currency": r.commission_ccy,
        }
    elif r.commission_model == "per_trade":
        assert r.commission_amount is not None
        commission = {
            "model": "per_order",
            "amount": r.commission_amount / cfg.commission.per_trade_sides,
            "currency": r.commission_ccy,
        }
    else:
        raise ConfigError(f"{r.broker_symbol}: no profile rule for {r.commission_model!r}")
    swap: dict[str, Any]
    if r.swap_model == "none" and r.region == cfg.etf_region:
        e = cfg.swap.etf_assumed
        swap = {
            "model": "annual_rate",
            "long": e.long / 100.0,
            "short": e.short / 100.0,
            "day_count": cfg.swap.percent_day_count,
            **_rollover(cfg, e.triple_weekday),
        }
        to_verify.append("swap_assumed")
    elif r.swap_model == "none":
        swap = {"model": "none"}
    elif r.swap_model == "points":
        swap = {
            "model": "points_per_day",
            "long": r.swap_long,
            "short": r.swap_short,
            "point_size": r.point_size,
            **_rollover(cfg, r.triple_weekday),
        }
    elif r.swap_model == "currency":
        swap = {
            "model": "currency_per_lot_day",
            "long": r.swap_long,
            "short": r.swap_short,
            **_rollover(cfg, r.triple_weekday),
        }
        if r.quote_ccy != "USD":
            to_verify.append("swap_non_usd_index")
    else:
        assert r.swap_long is not None and r.swap_short is not None
        swap = {
            "model": "annual_rate",
            "long": r.swap_long / 100.0,
            "short": r.swap_short / 100.0,
            "day_count": cfg.swap.percent_day_count,
            **_rollover(cfg, r.triple_weekday),
        }
    if asset_class not in cfg.slippage:
        raise ConfigError(f"no slippage assumption for asset class {asset_class!r}")
    d.update(
        source_note=note,
        pip_size=r.point_size * cfg.pip_per_point if asset_class in cfg.pip_classes else None,
        spread=spread,
        commission=commission,
        swap=swap,
        slippage=cfg.slippage[asset_class].model_dump(),
        broker_symbol=r.broker_symbol,
        quote_ccy=r.quote_ccy,
        point_value=r.point_value,
        contract_size=r.contract_size,
        volume_step=r.volume_step_lots,
        min_volume=r.min_volume_lots,
        volume_step_assumed=False,
        to_verify=to_verify,
    )
    CostProfile.model_validate(d)  # fail early on a bad rule
    return d


def proxy_profile_dict(spec: Iterable[SpecRow], sha: str, cfg: MonetaConfig) -> dict[str, Any]:
    """D-324: the Moneta US-share model with the median bps of the broker's US shares."""
    rows = [r for r in spec if r.region == cfg.us_share_region and r.row_status == "ok"]
    if not rows:
        raise ConfigError(f"no ok rows for region {cfg.us_share_region!r}")
    bps = statistics.median(_bps(r) for r in rows)
    (sl, ss, triple), _ = Counter(
        (r.swap_long, r.swap_short, r.triple_weekday) for r in rows
    ).most_common(1)[0]
    (c_model,), _ = Counter((r.commission_model,) for r in rows).most_common(1)[0]
    if c_model != "none" or sl is None or ss is None:
        raise ConfigError(f"unexpected {cfg.us_share_region} model: commission {c_model}")
    d = {
        "name": cfg.proxy_profile,
        "status": "placeholder",
        "source_note": (
            f"D-324 proxy for research symbols not at the broker: Moneta {cfg.us_share_region} "
            f"model from spec sha256={sha} file date {cfg.source.file_date}; spread = median of "
            f"{len(rows)} broker US shares = {bps:.6g} bps; results are report-only (D-524)"
        ),
        "spread": {"mode": "fixed", "fixed": {"value": bps, "unit": "bps"}},
        "commission": {"model": "none"},
        "swap": {
            "model": "annual_rate",
            "long": sl / 100.0,
            "short": ss / 100.0,
            "day_count": cfg.swap.percent_day_count,
            **_rollover(cfg, triple),
        },
        "slippage": cfg.slippage[cfg.proxy_asset_class].model_dump(),
        "volume_step_assumed": True,  # D-314: 1 share per lot, step 1, minimum 1
    }
    CostProfile.model_validate(d)
    return d


# -- build -------------------------------------------------------------------------------------
@dataclass(frozen=True)
class BuildResult:
    profiles: list[dict[str, Any]]
    mapping: SymbolMapping
    to_verify: dict[str, list[str]]


def build_profiles(
    spec: Sequence[SpecRow],
    mapping: SymbolMapping,
    research: Mapping[str, str],
    sha: str,
    cfg: MonetaConfig,
) -> BuildResult:
    by_broker = {r.broker_symbol: r for r in spec}
    profiles = [proxy_profile_dict(spec, sha, cfg)]
    verify: dict[str, list[str]] = {}
    for m in mapping.mapped:
        d = profile_dict(by_broker[m.broker_symbol], research[m.research_symbol], sha, cfg)
        profiles.append(d)
        if d["to_verify"]:
            verify[m.research_symbol] = list(d["to_verify"])
    return BuildResult(profiles, mapping, verify)


PROFILES_HEADER = """\
# GENERATED by `sfac costs moneta build` from moneta_spec.csv, moneta.yaml and the symbol map.
# Never edit by hand: change the rules or the overrides and rebuild (T06b, D-520 ... D-526).
"""
ASSIGN_HEADER = """\
# GENERATED by `sfac costs moneta build`: broker-mapped research symbol -> Moneta profile.
# Merged with ../assignments.yaml (a symbol may not appear in both). Never edit by hand.
"""


def write_build(result: BuildResult, costs_dir: Path, review_csv: Path) -> None:
    mdir = costs_dir / MONETA_DIR
    body = yaml.safe_dump(
        {"profiles": result.profiles}, sort_keys=False, default_flow_style=None, width=10_000
    )
    (mdir / MONETA_PROFILES_FILE).write_bytes((PROFILES_HEADER + body).encode("utf-8"))
    symbols = {
        m.research_symbol: {"profile": profile_name(m.broker_symbol)} for m in result.mapping.mapped
    }
    body = yaml.safe_dump(
        {"symbols": symbols}, sort_keys=True, default_flow_style=None, width=10_000
    )
    (mdir / ASSIGNMENTS_FILE).write_bytes((ASSIGN_HEADER + body).encode("utf-8"))
    _write_rows(mdir / SYMBOL_MAP_CSV, [asdict(m) for m in result.mapping.mapped], MapRow)
    review_csv.parent.mkdir(parents=True, exist_ok=True)
    _write_rows(review_csv, [asdict(v) for v in result.mapping.review], ReviewRow)


def _write_rows(path: Path, rows: list[dict[str, Any]], cls: type) -> None:
    cols = [f.name for f in fields(cls)]
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(cols)
    for r in rows:
        w.writerow(
            [
                f"{r[c]:.4f}" if isinstance(r[c], float) and not math.isnan(r[c]) else _fmt(r[c])
                for c in cols
            ]
        )
    path.write_bytes(buf.getvalue().encode("utf-8"))


def read_symbol_map(costs_dir: Path) -> dict[str, str]:
    """research symbol -> broker symbol (empty when the map has not been built)."""
    path = costs_dir / MONETA_DIR / SYMBOL_MAP_CSV
    if not path.is_file():
        return {}
    return {r["research_symbol"]: r["broker_symbol"] for r in _read_csv(path)}


def build(
    costs_dir: Path, raw_root: Path, research: Mapping[str, str], review_csv: Path
) -> BuildResult:
    """Everything ``sfac costs moneta build`` does (see the module docstring)."""
    mdir = costs_dir / MONETA_DIR
    cfg, mcfg = load_moneta_config(mdir), load_mapping_config(mdir)
    spec = read_spec_csv(mdir / SPEC_CSV)
    sha = spec_sha256(mdir)
    names = read_names(raw_root / mcfg.names_file)
    mapping = map_symbols(
        spec,
        research,
        names,
        _read_csv(mdir / DUKASCOPY_MAP_CSV),
        _read_csv(mdir / OVERRIDES_CSV),
        mcfg,
    )
    if mcfg.prices_dir:  # D-341: the price check needs the raw daily files
        day = dt.date.fromisoformat(cfg.source.file_date)
        priced = with_price_check(mapping.review, spec, day, raw_root / mcfg.prices_dir)
        mapping = SymbolMapping(mapping.mapped, priced)
    result = build_profiles(spec, mapping, research, sha, cfg)
    write_build(result, costs_dir, review_csv)
    return result
